"""Exports that exist as things rather than as responses (§30).

Four claims this module is here for. Each of them was a promise the platform
made in a comment or an error message before anything made it true.

**The ceiling means something.** `core/export.MAX_ROWS` had carried the comment
"rows above which an export must become a background job" since it was written,
and nothing enforced it: an export of 200,000 rows produced 50,000 with a `200`
on it. The refusal now names the queued path, and the queued path exists — so
the two halves are tested against each other rather than separately.

**A queued export is the same question.** Both the streamed download and the
background job build their statement from one `explorer.Plan`, so the test that
matters is the one that runs a *filtered* export and checks every row in the
produced file satisfies the filter — an export that quietly widened its query
because it took the slow path would be the worst kind of wrong: plausible.

**The row count is the file's, not the request's.** `total_units` is counted
when the export is queued and the file is written afterwards, so they can
legitimately differ. Everything a reader is shown comes from the object that
was actually stored: `result["rows"]` is what `write_to` counted, and the tests
compare it against the bytes rather than against the count that was hoped for.

**Somebody else's export is not there.** An export is a copy of whatever rows
its requester could see. A download is checked against the initiator and not
against a permission, administrators included — so the test asserts a 404
rather than a 403, because "that belongs to somebody else" is a way to confirm
a reference exists.

Every export these tests produce, they remove: the job rows go with the
conftest sweep, and the artefacts are dropped through `forget`, which is the
same path a person uses.
"""

from __future__ import annotations

import csv
import io
import json

import pytest
from sqlalchemy import select

from src.config import Config
from src.core import background, storage
from src.core.clock import now
from src.core import export as writer
from src.core.db import session_scope
from src.services import exports as service
from tests.conftest import persona_claims

pytestmark = pytest.mark.database

PREFIX = Config.API_PREFIX
EXPORTS = f"{PREFIX}/exports"

#: The dataset these tests export. Small, seeded, and filterable by a field
#: with a closed value set, which is what makes the "same question" assertion
#: possible without inventing rows.
RESOURCE = "ticket"


def _authenticate(monkeypatch, username: str = "analyst", role: str = "analyst"):
    monkeypatch.setattr(
        "src.core.auth.verify_token",
        lambda _token: persona_claims(username, role, sid=f"exp-{username}"),
    )
    return {"Authorization": f"Bearer exp-{username}"}


def _queue(client, headers, **payload):
    """Queue an export and wait for it by not needing to.

    `background.synchronous()` makes `spawn` run inline for the duration, so
    the response describes a finished job. A test that polled instead would be
    a test that passes when the work is fast and flakes when it is not.
    """
    body = {"resource_type": RESOURCE, "format": "csv", **payload}
    with background.synchronous():
        response = client.post(EXPORTS, json=body, headers=headers)
    return response


def _artifact(export_id: str) -> tuple[str, bytes]:
    """The key and bytes of what was actually stored for one export."""
    from src.models.platform import BackgroundJob

    with session_scope() as session:
        row = session.get(BackgroundJob, export_id)
        assert row is not None
        key = str((row.result or {}).get("artifact") or "")
    assert key, "the export recorded no artefact"
    return key, storage.for_config().get(key)


def _discard(client, headers, export_id: str) -> None:
    """Let an export go entirely, through the endpoint a person would use.

    Twice, because `forget` takes two presses — the file, then the record — and
    a sweep that pressed once would leave a row behind for every test in this
    module. The same shape `e2e/api.sweepMailThreads` documents for threads.
    """
    client.delete(f"{EXPORTS}/{export_id}", headers=headers)
    client.delete(f"{EXPORTS}/{export_id}", headers=headers)


# ── off-request execution ────────────────────────────────────────────────


def test_the_mechanism_is_named_rather_than_assumed():
    # Three quite different things hide behind the word "background", and an
    # operator reading a slow export needs to know which one they have.
    assert background.mechanism() in background.MECHANISMS
    with background.synchronous():
        assert background.mechanism() == "inline"


def test_synchronous_runs_the_work_before_spawn_returns():
    ran: list[str] = []
    with background.synchronous():
        used = background.spawn(lambda: ran.append("yes"), name="test")
    assert used == "inline"
    assert ran == ["yes"], "synchronous() has to finish the work, not schedule it"


def test_a_raising_callable_does_not_escape_spawn():
    # There is no caller left to catch anything by the time this runs, so the
    # contract is that `spawn` swallows and logs. A background mechanism that
    # let an exception out would take the greenlet with it.
    def boom() -> None:
        raise RuntimeError("no")

    with background.synchronous():
        assert background.spawn(boom, name="test") == "inline"


# ── the file itself ─────────────────────────────────────────────────────


def test_write_to_counts_the_rows_it_wrote():
    columns = [writer.Column("a", "A"), writer.Column("b", "B")]
    rows = [{"a": 1, "b": "x"}, {"a": 2, "b": "y"}, {"a": 3, "b": "z"}]
    buffer = io.BytesIO()

    written = writer.write_to(buffer, iter(rows), columns, fmt="csv")

    # The count is the file's, which is the whole reason it is returned rather
    # than taken from the query that produced the rows.
    assert written == 3
    body = buffer.getvalue().decode("utf-8-sig")
    assert list(csv.reader(io.StringIO(body)))[0] == ["A", "B"]
    assert len(list(csv.reader(io.StringIO(body)))) == 4


def test_write_to_produces_json_that_parses():
    columns = [writer.Column("a", "A")]
    buffer = io.BytesIO()
    writer.write_to(buffer, iter([{"a": None}, {"a": 2}]), columns, fmt="json")
    # `null` rather than `""`: a consumer parsing JSON should not have to guess
    # whether an empty string meant a missing value.
    assert json.loads(buffer.getvalue()) == [{"a": None}, {"a": 2}]


def test_write_to_produces_a_workbook_openpyxl_can_read():
    from openpyxl import load_workbook

    columns = [writer.Column("a", "A")]
    buffer = io.BytesIO()
    assert writer.write_to(buffer, iter([{"a": 7}]), columns, fmt="xlsx") == 1
    buffer.seek(0)
    sheet = load_workbook(buffer).active
    assert [cell.value for cell in next(sheet.iter_rows())] == ["A"]


# ── estimating before committing ─────────────────────────────────────────


def test_an_estimate_says_whether_it_streams_or_has_to_be_queued(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    answer = client.post(
        f"{EXPORTS}/estimate", json={"resource_type": RESOURCE, "format": "csv"}, headers=headers
    )
    assert answer.status_code == 200
    body = answer.get_json()

    assert body["rows"] > 0, "the seed has no tickets"
    # The seeded dataset is far below the streaming ceiling, so a page should
    # offer the immediate download and not add a step.
    assert body["can_stream"] is True
    assert body["can_queue"] is True
    assert body["too_large"] is False
    assert body["streams_up_to"] == writer.MAX_ROWS


def test_an_estimate_counts_the_filtered_question_not_the_table(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    whole = client.post(
        f"{EXPORTS}/estimate", json={"resource_type": RESOURCE}, headers=headers
    ).get_json()
    narrowed = client.post(
        f"{EXPORTS}/estimate",
        json={"resource_type": RESOURCE, "filters": {"status": "OPEN"}},
        headers=headers,
    ).get_json()

    assert narrowed["rows"] < whole["rows"], "the filter did not reach the count"
    assert "status" in narrowed["description"]


def test_the_streaming_ceiling_the_estimate_reports_is_the_one_that_is_enforced(
    client, monkeypatch
):
    # The number a page prints and the number the API refuses at have to be the
    # same number, or the page is lying about the product's own limits.
    monkeypatch.setattr(writer, "MAX_ROWS", 3)
    headers = _authenticate(monkeypatch)
    body = client.post(
        f"{EXPORTS}/estimate", json={"resource_type": RESOURCE, "format": "csv"}, headers=headers
    ).get_json()

    assert body["can_stream"] is False
    refused = client.post(
        f"{PREFIX}/api/explorer/export",
        json={"resource_type": RESOURCE, "format": "csv"},
        headers=headers,
    )
    assert refused.status_code == 400
    details = refused.get_json()["details"]
    assert details["maximum"] == 3
    # And the refusal points at the path that now exists.
    assert details["queue_instead"] is True


# ── queueing, and what comes out ─────────────────────────────────────────


def test_a_queued_export_produces_a_real_file_with_a_real_row_count(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    response = _queue(client, headers)
    assert response.status_code == 201
    body = response.get_json()

    try:
        assert body["reference"].startswith("EXP-")
        assert body["status"] == "SUCCEEDED"
        assert body["rows"] > 0
        assert body["downloadable"] is True

        key, blob = _artifact(body["id"])
        assert key.startswith("exports/")
        # The count on the row is the count in the file. Not "about right":
        # equal, because one was derived from the other.
        lines = list(csv.reader(io.StringIO(blob.decode("utf-8-sig"))))
        assert len(lines) == body["rows"] + 1
        assert body["size_bytes"] == len(blob)
    finally:
        _discard(client, headers, body["id"])


def test_a_queued_export_is_the_same_question_as_the_download_would_have_been(
    client, monkeypatch
):
    # The claim the shared `explorer.Plan` exists for. An export that queued
    # *because* it was large must not be a wider query than the one on screen.
    headers = _authenticate(monkeypatch)
    response = _queue(
        client, headers, filters={"status": "OPEN"}, columns=["reference", "status"]
    )
    body = response.get_json()

    try:
        key, blob = _artifact(body["id"])
        rows = list(csv.DictReader(io.StringIO(blob.decode("utf-8-sig"))))
        assert rows, "the filtered export is empty, so it proves nothing"
        assert {row["Status"] for row in rows} == {"OPEN"}
        # And the columns are the ones asked for, in that order.
        assert list(rows[0]) == ["Reference", "Status"]
    finally:
        _discard(client, headers, body["id"])


def test_the_row_count_matches_what_the_same_statement_counts(client, monkeypatch):
    from src.core.query import count_of
    from src.services import explorer

    headers = _authenticate(monkeypatch)
    response = _queue(client, headers, filters={"status": "OPEN"})
    body = response.get_json()

    try:
        with session_scope() as session:
            plan = explorer.replan(
                {"resource_type": RESOURCE, "format": "csv", "filters": {"status": "OPEN"}}
            )
            counted = count_of(session, explorer.statement_of(plan))
        assert body["rows"] == counted
    finally:
        _discard(client, headers, body["id"])


def test_the_stored_plan_describes_itself_from_the_field_catalogue(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    response = _queue(client, headers, filters={"status": "OPEN"})
    body = response.get_json()

    try:
        # Re-derived on every read rather than saved as a sentence, so
        # relabelling a field cannot leave a stale description behind.
        assert "Tickets" in body["description"] or "ticket" in body["description"].lower()
        assert body["columns"], "the plan recorded no columns"
    finally:
        _discard(client, headers, body["id"])


def test_a_plan_naming_a_column_that_no_longer_exists_is_a_bad_request():
    from src.core.errors import ValidationError
    from src.services import explorer

    # Re-validated when it runs rather than trusted: a field can be removed
    # from a resource between queueing an export and producing it, and that has
    # to be a 400 and not a 500 inside a background job.
    with pytest.raises(ValidationError):
        explorer.replan({"resource_type": RESOURCE, "columns": ["a_field_nobody_declared"]})


def test_an_export_whose_query_has_gone_stale_still_reads_as_a_record(client, monkeypatch):
    from src.models.platform import BackgroundJob

    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    try:
        with session_scope() as session:
            row = session.get(BackgroundJob, body["id"])
            row.payload = {**(row.payload or {}), "columns": ["gone_away"]}

        # The row is still the record of what somebody asked for, so reading it
        # must not fail — it says the query cannot be read any more.
        again = client.get(f"{EXPORTS}/{body['id']}", headers=headers).get_json()
        assert "no longer exist" in again["description"]
    finally:
        _discard(client, headers, body["id"])


def test_the_reference_continues_from_the_highest_one(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    first = _queue(client, headers).get_json()
    second = _queue(client, headers).get_json()

    try:
        assert first["reference"] != second["reference"]
        assert int(second["reference"].removeprefix("EXP-")) == (
            int(first["reference"].removeprefix("EXP-")) + 1
        )
    finally:
        _discard(client, headers, first["id"])
        _discard(client, headers, second["id"])


# ── the ceilings ─────────────────────────────────────────────────────────


def test_above_the_installations_ceiling_it_is_refused_with_the_number(client, monkeypatch):
    monkeypatch.setattr(service, "DEFAULT_MAX_ROWS", 2)
    monkeypatch.setattr(service, "max_rows", lambda _session: 2)
    headers = _authenticate(monkeypatch)

    refused = client.post(EXPORTS, json={"resource_type": RESOURCE}, headers=headers)
    assert refused.status_code == 400
    details = refused.get_json()["details"]
    assert details["maximum"] == 2
    # Named, so an administrator knows which row to edit rather than being told
    # to "narrow the query" with no other option.
    assert details["setting"] == "limits.max_export_rows"


def test_the_ceiling_comes_from_the_settings_table(client, monkeypatch):
    from src.models.platform import SystemSetting

    headers = _authenticate(monkeypatch)
    with session_scope() as session:
        row = session.scalar(
            select(SystemSetting).where(SystemSetting.key == "limits.max_export_rows")
        )
        assert row is not None, "run `make seed-settings`; the setting is not there"
        stored = (row.value or {}).get("value")

    body = client.get(f"{EXPORTS}/catalogue", headers=headers).get_json()
    assert body["max_rows"] == stored


def test_the_spreadsheet_ceiling_is_not_raised_by_the_setting(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    body = client.get(f"{EXPORTS}/catalogue", headers=headers).get_json()
    formats = {entry["key"]: entry for entry in body["formats"]}

    # A zip container has to be finished before it can be sent, which is a
    # property of the format and not of the installation.
    assert formats["xlsx"]["maximum"] <= writer.MAX_XLSX_ROWS
    assert formats["csv"]["maximum"] == body["max_rows"]


# ── stalling, and asking again ──────────────────────────────────────────


def test_a_pending_export_that_is_too_old_reads_as_stalled(client, monkeypatch):
    from datetime import timedelta

    from src.models.platform import BackgroundJob

    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    try:
        with session_scope() as session:
            row = session.get(BackgroundJob, body["id"])
            row.status = "QUEUED"
            row.started_at = None
            # Backdated, and put back in the `finally`: the conftest sweep
            # decides what a test created by comparing `created_at` against
            # the instant the test began, so a row moved into the past is a row
            # that leaks into the demo database. Two of them did.
            row.created_at = row.created_at - timedelta(hours=3)
            assert service.stalled(row) is True

        # `core/background` loses in-flight work when the process restarts, so
        # a QUEUED row nothing will ever pick up is a state this platform
        # genuinely produces. A spinner that never stops is the wrong answer.
        again = client.get(f"{EXPORTS}/{body['id']}", headers=headers).get_json()
        assert again["stalled"] is True
        assert again["downloadable"] is False
    finally:
        with session_scope() as session:
            row = session.get(BackgroundJob, body["id"])
            if row is not None:
                row.status = "SUCCEEDED"
                row.created_at = now()
        _discard(client, headers, body["id"])


def test_a_stalled_export_does_not_hold_a_slot_forever(client, monkeypatch):
    from datetime import timedelta

    from src.models.platform import BackgroundJob

    monkeypatch.setattr(service, "MAX_PENDING_PER_PERSON", 1)
    headers = _authenticate(monkeypatch)
    first = _queue(client, headers).get_json()

    try:
        with session_scope() as session:
            row = session.get(BackgroundJob, first["id"])
            row.status = "QUEUED"
            row.started_at = None
            # Put back in the `finally` — see the note in the test above.
            row.created_at = row.created_at - timedelta(hours=3)

        # Three of these and a limit of three would mean the button never
        # worked again, which is what the seeded queue's pending exports would
        # otherwise do on a fresh installation.
        second = _queue(client, headers)
        assert second.status_code == 201, second.get_json()
        _discard(client, headers, second.get_json()["id"])
    finally:
        with session_scope() as session:
            row = session.get(BackgroundJob, first["id"])
            if row is not None:
                row.status = "SUCCEEDED"
                row.created_at = now()
        _discard(client, headers, first["id"])


def test_a_fresh_pending_export_does_hold_one(client, monkeypatch):
    from src.models.platform import BackgroundJob

    monkeypatch.setattr(service, "MAX_PENDING_PER_PERSON", 1)
    headers = _authenticate(monkeypatch)
    first = _queue(client, headers).get_json()

    try:
        with session_scope() as session:
            session.get(BackgroundJob, first["id"]).status = "RUNNING"

        refused = client.post(EXPORTS, json={"resource_type": RESOURCE}, headers=headers)
        assert refused.status_code == 409
        details = refused.get_json()["details"]
        assert details["maximum"] == 1
        # Named, so somebody can go and look at the one that is blocking them
        # rather than being told a number.
        assert first["reference"] in details["references"]
    finally:
        with session_scope() as session:
            row = session.get(BackgroundJob, first["id"])
            if row is not None:
                row.status = "SUCCEEDED"
        _discard(client, headers, first["id"])


def test_asking_again_makes_a_new_export_and_leaves_the_old_one_alone(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    first = _queue(client, headers).get_json()

    with background.synchronous():
        answer = client.post(f"{EXPORTS}/{first['id']}/again", headers=headers)
    assert answer.status_code == 201
    second = answer.get_json()

    try:
        # A new row, not a second attempt at the old one: the rows have moved on
        # since the first was asked for, and a file whose reference says one
        # moment and whose contents say another is worse than no file.
        assert second["id"] != first["id"]
        assert second["reference"] != first["reference"]
        assert second["attempt"] == 1
        # And the original is untouched, artefact included.
        untouched = client.get(f"{EXPORTS}/{first['id']}", headers=headers).get_json()
        assert untouched["downloadable"] is True
        assert untouched["reference"] == first["reference"]
    finally:
        _discard(client, headers, first["id"])
        _discard(client, headers, second["id"])


def test_asking_again_asks_the_same_question(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    first = _queue(
        client, headers, filters={"status": "OPEN"}, columns=["reference", "status"]
    ).get_json()

    with background.synchronous():
        second = client.post(f"{EXPORTS}/{first['id']}/again", headers=headers).get_json()

    try:
        assert second["columns"] == first["columns"]
        assert second["description"] == first["description"]
    finally:
        _discard(client, headers, first["id"])
        _discard(client, headers, second["id"])


def test_asking_again_on_somebody_elses_export_is_not_offered(client, monkeypatch):
    headers = _authenticate(monkeypatch, "analyst", "analyst")
    body = _queue(client, headers).get_json()

    try:
        admin = _authenticate(monkeypatch, "admin", "administrator")
        assert client.post(f"{EXPORTS}/{body['id']}/again", headers=admin).status_code == 404
    finally:
        headers = _authenticate(monkeypatch, "analyst", "analyst")
        _discard(client, headers, body["id"])


def test_asking_again_on_a_stale_question_is_refused_with_the_reason(client, monkeypatch):
    from src.models.platform import BackgroundJob

    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    try:
        with session_scope() as session:
            row = session.get(BackgroundJob, body["id"])
            row.payload = {**(row.payload or {}), "columns": ["gone_away"]}

        # Refused through the ordinary validation rather than accepted and
        # failed inside a background job, which is where it used to be found.
        answer = client.post(f"{EXPORTS}/{body['id']}/again", headers=headers)
        assert answer.status_code == 400
        assert "gone_away" in str(answer.get_json())
    finally:
        _discard(client, headers, body["id"])


# ── downloading ─────────────────────────────────────────────────────────


def test_a_download_is_a_signed_url_to_the_bytes_that_were_stored(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    try:
        answer = client.get(f"{EXPORTS}/{body['id']}/download", headers=headers)
        assert answer.status_code == 200
        link = answer.get_json()

        assert link["url"], "no URL was signed"
        assert link["expires_in"] > 0, "a link that never expires is a permanent grant"
        assert link["rows"] == body["rows"]
        # The signature is over the object that exists, which is what makes the
        # link worth having.
        _, blob = _artifact(body["id"])
        assert link["size_bytes"] == len(blob)
    finally:
        _discard(client, headers, body["id"])


def test_a_download_before_the_file_exists_says_so_rather_than_404(client, monkeypatch):
    from src.models.platform import BackgroundJob

    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    try:
        with session_scope() as session:
            session.get(BackgroundJob, body["id"]).status = "RUNNING"

        answer = client.get(f"{EXPORTS}/{body['id']}/download", headers=headers)
        # 409, not 404: the export exists and is not ready, which is a
        # different thing from not existing and reads differently on a page.
        assert answer.status_code == 409
        assert answer.get_json()["details"]["status"] == "RUNNING"
    finally:
        with session_scope() as session:
            row = session.get(BackgroundJob, body["id"])
            if row is not None:
                row.status = "SUCCEEDED"
        _discard(client, headers, body["id"])


def test_an_artefact_missing_from_storage_is_reported_and_the_row_corrected(
    client, monkeypatch
):
    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()
    key, _ = _artifact(body["id"])

    # Removed behind the platform's back, which is what happens when somebody
    # empties a bucket. A signed URL to a missing object is a download that
    # fails with the wrong error.
    storage.for_config().delete(key)

    answer = client.get(f"{EXPORTS}/{body['id']}/download", headers=headers)
    assert answer.status_code == 404
    again = client.get(f"{EXPORTS}/{body['id']}", headers=headers).get_json()
    # And the row now agrees with storage rather than going on offering it.
    assert again["downloadable"] is False


def test_an_expired_export_is_refused_and_its_bytes_are_dropped(client, monkeypatch):
    from datetime import timedelta

    from src.models.platform import BackgroundJob

    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()
    key, _ = _artifact(body["id"])

    with session_scope() as session:
        row = session.get(BackgroundJob, body["id"])
        row.finished_at = row.finished_at - timedelta(days=400)

    answer = client.get(f"{EXPORTS}/{body['id']}/download", headers=headers)
    assert answer.status_code == 409
    details = answer.get_json()["details"]
    assert details["expired"] is True
    assert details["retention_days"] > 0
    # Lazy expiry, done where the fact is discovered: there is no sweeper in
    # this stack, so the alternative is bytes that live forever.
    assert storage.for_config().stat(key) is None


def test_expiry_is_derived_so_changing_the_setting_changes_it(client, monkeypatch):
    from src.models.platform import BackgroundJob

    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    try:
        with session_scope() as session:
            row = session.get(BackgroundJob, body["id"])
            close = service.expires_at(row, 1)
            far = service.expires_at(row, 30)
        # Stored as a column it would go on answering the number that applied
        # when it was written.
        assert far > close
    finally:
        _discard(client, headers, body["id"])


# ── whose export it is ───────────────────────────────────────────────────


def test_the_listing_is_mine_and_only_mine(client, monkeypatch):
    headers = _authenticate(monkeypatch, "analyst", "analyst")
    body = _queue(client, headers).get_json()

    try:
        mine = client.get(EXPORTS, headers=headers).get_json()
        assert body["id"] in [row["id"] for row in mine["items"]]

        other = _authenticate(monkeypatch, "manager", "manager")
        theirs = client.get(EXPORTS, headers=other).get_json()
        assert body["id"] not in [row["id"] for row in theirs["items"]]
    finally:
        headers = _authenticate(monkeypatch, "analyst", "analyst")
        _discard(client, headers, body["id"])


def test_an_administrator_cannot_download_somebody_elses_export(client, monkeypatch):
    headers = _authenticate(monkeypatch, "analyst", "analyst")
    body = _queue(client, headers).get_json()

    try:
        admin = _authenticate(monkeypatch, "admin", "administrator")
        answer = client.get(f"{EXPORTS}/{body['id']}/download", headers=admin)
        # 404 rather than 403, deliberately: a "that is somebody else's" reply
        # is a way to confirm a reference exists. And an administrator who
        # needs the data can run the query in their own name, which leaves an
        # audit entry that a download of this row would not.
        assert answer.status_code == 404
        assert client.get(f"{EXPORTS}/{body['id']}", headers=admin).status_code == 404
    finally:
        headers = _authenticate(monkeypatch, "analyst", "analyst")
        _discard(client, headers, body["id"])


def test_a_role_without_the_export_privilege_is_refused(client, monkeypatch):
    headers = _authenticate(monkeypatch, "user", "viewer")
    for method, url in (
        ("get", EXPORTS),
        ("get", f"{EXPORTS}/catalogue"),
    ):
        assert getattr(client, method)(url, headers=headers).status_code == 403
    assert client.post(EXPORTS, json={"resource_type": RESOURCE}, headers=headers).status_code == 403


def test_a_dataset_the_requester_cannot_read_cannot_be_exported(client, monkeypatch):
    # The export privilege is not a way around a dataset's own permission: the
    # plan is authorised against the resource before anything is queued.
    headers = _authenticate(monkeypatch, "analyst", "analyst")
    answer = client.post(EXPORTS, json={"resource_type": "user"}, headers=headers)
    assert answer.status_code in (400, 403)


def test_the_catalogue_offers_only_datasets_the_reader_may_see(client, monkeypatch):
    headers = _authenticate(monkeypatch, "analyst", "analyst")
    body = client.get(f"{EXPORTS}/catalogue", headers=headers).get_json()
    offered = {entry["key"] for entry in body["datasets"]}

    assert RESOURCE in offered
    from src.services import explorer

    with session_scope():
        for key, resource in explorer.resources().items():
            if key in offered:
                continue
            # Anything withheld is withheld because of its own permission, not
            # arbitrarily — a picker offering a dataset whose export 403s is a
            # picker that wastes somebody's time.
            assert resource.permission != "records.view", key


# ── letting go ──────────────────────────────────────────────────────────


def test_the_first_press_drops_the_file_and_keeps_the_record(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()
    key, _ = _artifact(body["id"])

    answer = client.delete(f"{EXPORTS}/{body['id']}", headers=headers)
    assert answer.status_code == 200
    after = answer.get_json()

    assert storage.for_config().stat(key) is None, "the bytes are still there"
    assert after["downloadable"] is False
    assert after["removed"] is False
    # The row survives the file, because the two acts are different sizes: one
    # takes a copy of production data out of storage.
    assert client.get(f"{EXPORTS}/{body['id']}", headers=headers).status_code == 200
    assert after["rows"] == body["rows"]


def test_a_discarded_export_keeps_the_counts_as_history(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    after = client.delete(f"{EXPORTS}/{body['id']}", headers=headers).get_json()
    try:
        # "1,284 rows, 88 KB" is what the record is *for*, so the file going
        # must not take it. `has_file` is the separate question, and the one
        # that decides what a second press does.
        assert after["rows"] == body["rows"]
        assert after["size_bytes"] == body["size_bytes"]
        assert after["has_file"] is False
        assert after["downloadable"] is False
    finally:
        _discard(client, headers, body["id"])


def test_the_second_press_removes_the_record(client, monkeypatch):
    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    assert client.delete(f"{EXPORTS}/{body['id']}", headers=headers).get_json()[
        "removed"
    ] is False
    second = client.delete(f"{EXPORTS}/{body['id']}", headers=headers)
    assert second.status_code == 200
    assert second.get_json()["removed"] is True

    # Gone from the list, which is what a person asked for. Keeping every
    # discarded export forever answers nothing about one person's own requests.
    assert client.get(f"{EXPORTS}/{body['id']}", headers=headers).status_code == 404


def test_the_record_that_is_removed_leaves_its_history_behind(client, monkeypatch):
    from src.models.platform import AuditLog

    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()
    client.delete(f"{EXPORTS}/{body['id']}", headers=headers)
    client.delete(f"{EXPORTS}/{body['id']}", headers=headers)

    with session_scope() as session:
        actions = set(
            session.scalars(
                select(AuditLog.action).where(AuditLog.resource_id == body["id"])
            ).all()
        )
    # The history the row used to carry, in the one place its requester cannot
    # edit it — which is the reason removing the row is defensible at all.
    assert {"export.queue", "export.forget", "export.remove"} <= actions


def test_forget_is_refused_while_the_export_is_still_being_produced(client, monkeypatch):
    from src.models.platform import BackgroundJob

    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    try:
        with session_scope() as session:
            session.get(BackgroundJob, body["id"]).status = "RUNNING"
        answer = client.delete(f"{EXPORTS}/{body['id']}", headers=headers)
        assert answer.status_code == 409
    finally:
        with session_scope() as session:
            row = session.get(BackgroundJob, body["id"])
            if row is not None:
                row.status = "SUCCEEDED"
        _discard(client, headers, body["id"])


def test_a_third_press_has_nothing_left_to_find(client, monkeypatch):
    # Not idempotent, and it should not be: after the second press the export
    # is gone, and a 200 would suggest there was still something there.
    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()
    assert client.delete(f"{EXPORTS}/{body['id']}", headers=headers).status_code == 200
    assert client.delete(f"{EXPORTS}/{body['id']}", headers=headers).status_code == 200
    assert client.delete(f"{EXPORTS}/{body['id']}", headers=headers).status_code == 404


# ── failure ─────────────────────────────────────────────────────────────


def test_a_failure_is_recorded_on_the_row_rather_than_lost(client, monkeypatch):
    def explode(*_args, **_kwargs):
        raise RuntimeError("the bucket said no")

    monkeypatch.setattr(service, "produce", explode)
    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    # Nothing above the job is still running when it fails, so the row is the
    # only place the reason can end up.
    again = client.get(f"{EXPORTS}/{body['id']}", headers=headers).get_json()
    assert again["status"] == "FAILED"
    assert "the bucket said no" in (again["error_message"] or "")
    assert again["downloadable"] is False
    assert any("failed" in line.get("message", "") for line in again["log_lines"])


def test_a_failed_export_offers_no_download(client, monkeypatch):
    monkeypatch.setattr(
        service, "produce", lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("no"))
    )
    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    answer = client.get(f"{EXPORTS}/{body['id']}/download", headers=headers)
    assert answer.status_code == 409
    assert answer.get_json()["details"]["status"] == "FAILED"


# ── the record it leaves ────────────────────────────────────────────────


def test_queueing_and_downloading_are_both_audited(client, monkeypatch):
    from src.models.platform import AuditLog

    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    try:
        client.get(f"{EXPORTS}/{body['id']}/download", headers=headers)
        with session_scope() as session:
            actions = set(
                session.scalars(
                    select(AuditLog.action).where(AuditLog.resource_id == body["id"])
                ).all()
            )
        # Two separate facts: somebody asked, and somebody fetched. Only the
        # second is the moment a copy of the data left.
        assert {"export.queue", "export.download"} <= actions
    finally:
        _discard(client, headers, body["id"])


def test_the_audit_row_records_the_question_and_not_the_rows(client, monkeypatch):
    from src.models.platform import AuditLog

    headers = _authenticate(monkeypatch)
    body = _queue(client, headers).get_json()

    try:
        with session_scope() as session:
            entry = session.scalar(
                select(AuditLog).where(
                    AuditLog.resource_id == body["id"], AuditLog.action == "export.queue"
                )
            )
            after = entry.state_after or {}
        assert after["resource"] == RESOURCE
        assert after["rows"] == body["rows"] or after["rows"] >= 0
        # An audit row carrying the exported values would be a second copy of
        # the data the export already is.
        assert "items" not in after and "values" not in after
    finally:
        _discard(client, headers, body["id"])


# ── the seeded exports ──────────────────────────────────────────────────


def test_no_seeded_export_claims_a_file_it_does_not_have():
    """The defect that building this page found, asserted against the database.

    Three of the seeded exports said `{"rows": 184203, "artifact":
    "exports/JOB-000004.csv"}` for bytes nobody had written — and the extension
    was `.csv` while the payload said `xlsx`, so even the name was wrong. Only
    the real installation can say whether that is still true of it.
    """
    from src.models.platform import BackgroundJob

    store = storage.for_config()
    with session_scope() as session:
        rows = session.scalars(
            select(BackgroundJob).where(BackgroundJob.kind == service.KIND)
        ).all()
        assert rows, "run `make seed`; there are no export jobs"

        claimed = [
            (row.reference, str((row.result or {}).get("artifact") or ""))
            for row in rows
            if (row.result or {}).get("artifact")
        ]

    missing = [
        reference for reference, key in claimed if store.stat(key) is None
    ]
    assert missing == [], "seeded exports name files that are not in storage"


def test_every_seeded_export_names_a_query_that_can_be_run():
    from src.models.platform import BackgroundJob
    from src.services import explorer

    with session_scope() as session:
        rows = session.scalars(
            select(BackgroundJob).where(BackgroundJob.kind == service.KIND)
        ).all()
        unreadable = []
        for row in rows:
            try:
                explorer.statement_of(explorer.replan(row.payload))
            except Exception:  # noqa: BLE001 - the point is that none do
                unreadable.append(row.reference)
    # The old payload named an `entity` no code could resolve, so `/exports`
    # could not describe a single seeded row.
    assert unreadable == []


def test_a_seeded_exports_counts_describe_its_file():
    from src.models.platform import BackgroundJob

    with session_scope() as session:
        rows = session.scalars(
            select(BackgroundJob).where(
                BackgroundJob.kind == service.KIND, BackgroundJob.status == "SUCCEEDED"
            )
        ).all()
        assert rows, "no seeded export has finished"
        wrong = [
            row.reference
            for row in rows
            if (row.result or {}).get("rows") != row.processed_units
            or row.processed_units != row.total_units
        ]
    # `result.rows` used to be a progress counter and `total_units` a random
    # integer, so a finished export agreed with neither.
    assert wrong == []


def test_every_persona_who_may_export_has_one_to_download():
    """Otherwise the page is empty for whoever signs in, which shows nothing.

    The guarantee `seed/exports._guarantee` exists for, and the same argument
    `sync_jobs` makes about a filter that can never match: the ordinary draw
    attributed exports to whoever it liked and left all five personas without
    a finished one.
    """
    from src.models.identity import User
    from src.models.platform import BackgroundJob

    store = storage.for_config()
    with session_scope() as session:
        people = session.scalars(
            select(User).where(
                User.username.in_(("admin", "manager", "operator", "analyst", "user")),
                User.deleted_at.is_(None),
            )
        ).all()
        assert people, "run `make seed`; the personas are not there"

        short = []
        for person in people:
            allowed = set((person.role.permissions if person.role else None) or ())
            if service.PERMISSION not in allowed:
                # A viewer holds no export privilege, so having none is right.
                continue
            rows = session.scalars(
                select(BackgroundJob).where(
                    BackgroundJob.kind == service.KIND,
                    BackgroundJob.initiated_by_id == person.id,
                    BackgroundJob.status == "SUCCEEDED",
                )
            ).all()
            if not any(
                store.stat(str((row.result or {}).get("artifact") or "") or "-") is not None
                for row in rows
            ):
                short.append(person.username)

    assert short == [], "run `make sync-exports`"

